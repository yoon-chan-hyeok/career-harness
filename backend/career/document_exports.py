"""Physical PDF pages and editable DOCX, generated from one owned artifact version."""

import io
import re
from dataclasses import dataclass
from pathlib import Path
from threading import Lock
from typing import cast
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.styles.style import ParagraphStyle as WordParagraphStyle
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer

FONT_PATH = Path(__file__).with_name("assets") / "fonts"
FONT_LOCK = Lock()
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@dataclass(frozen=True)
class RenderedDocument:
    content: bytes
    media_type: str
    pages: int | None
    page_budget: int | str
    overflow: bool


def page_budget(data):
    value = data.get("page_budget")
    if value in (1, 2, "MASTER"):
        return value
    return {3500: 1, 7000: 2, 20000: "MASTER"}.get(data.get("character_budget"), 1)


def content_blocks(data):
    """Keep text lossless; a legacy document is one unaltered prose block."""
    if data.get("sections"):
        blocks = []
        for section in data["sections"]:
            blocks.append(("heading", str(section.get("title", ""))))
            for bullet in section.get("bullets", []):
                blocks.append(("bullet", str(bullet.get("text", ""))))
    else:
        blocks = [("text", str(data.get("text", "")))]
    if not any(text.strip() for _, text in blocks):
        raise ValueError("출력할 본문이 없습니다.")
    if sum(len(text) for _, text in blocks) > 100000:
        raise ValueError("출력 문서의 길이를 줄여 주세요.")
    if any(re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", text) for _, text in blocks):
        raise ValueError("문서의 제어 문자를 제거한 뒤 출력해 주세요.")
    return blocks


def _fonts():
    with FONT_LOCK:
        if "CareerKR" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("CareerKR", str(FONT_PATH / "CareerKR-Regular.ttf")))
            pdfmetrics.registerFont(TTFont("CareerKR-Bold", str(FONT_PATH / "CareerKR-Bold.ttf")))
            pdfmetrics.registerFontFamily("CareerKR", normal="CareerKR", bold="CareerKR-Bold")


def render_pdf(data) -> RenderedDocument:
    _fonts()
    blocks = content_blocks(data)
    budget = page_budget(data)
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, leftMargin=19 * mm, rightMargin=19 * mm,
                            topMargin=18 * mm, bottomMargin=18 * mm, title="이력서", author="")
    base = ParagraphStyle("CareerBody", parent=getSampleStyleSheet()["Normal"], fontName="CareerKR",
                          fontSize=10, leading=15, wordWrap="CJK", spaceAfter=5, alignment=TA_LEFT)
    heading = ParagraphStyle("CareerHeading", parent=base, fontName="CareerKR-Bold", fontSize=12,
                             leading=18, spaceBefore=10, spaceAfter=7, textColor=colors.HexColor("#284A40"))
    bullet = ParagraphStyle("CareerBullet", parent=base, leftIndent=10, firstLineIndent=-10)
    story = []
    for kind, text in blocks:
        markup = escape(text).replace("\n", "<br/>")
        if kind == "heading":
            story.append(KeepTogether([Spacer(1, 2), Paragraph(markup, heading)]))
        else:
            story.append(Paragraph(("• " if kind == "bullet" else "") + markup,
                                   bullet if kind == "bullet" else base))

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("CareerKR", 8)
        canvas.setFillColor(colors.HexColor("#6F7773"))
        canvas.drawRightString(A4[0] - 19 * mm, 10 * mm, str(document.page))
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    content = output.getvalue()
    pages = len(PdfReader(io.BytesIO(content)).pages)
    return RenderedDocument(content, "application/pdf", pages, budget,
                            isinstance(budget, int) and pages > budget)


def render_docx(data) -> RenderedDocument:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(19)
    section.top_margin = section.bottom_margin = Mm(18)
    normal = cast(WordParagraphStyle, document.styles["Normal"])
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Malgun Gothic")
    normal.paragraph_format.line_spacing = 1.2
    normal.paragraph_format.space_after = Pt(5)
    title_style = cast(WordParagraphStyle, document.styles["Heading 1"])
    title_style.font.name = "Calibri"
    title_style.font.size = Pt(12)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor.from_string("284A40")
    title_style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Malgun Gothic")
    title_style.paragraph_format.keep_with_next = True
    title_style.paragraph_format.space_before = Pt(10)
    title_style.paragraph_format.space_after = Pt(7)
    for kind, text in content_blocks(data):
        if kind == "heading":
            document.add_paragraph(text, style="Heading 1")
        else:
            paragraph = document.add_paragraph(("• " if kind == "bullet" else "") + text)
            paragraph.paragraph_format.widow_control = True
            if kind == "bullet":
                paragraph.paragraph_format.left_indent = Mm(3.5)
                paragraph.paragraph_format.first_line_indent = Mm(-3.5)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    output = io.BytesIO()
    document.save(output)
    # Word pagination depends on its renderer/fonts. Do not invent a DOCX page count.
    return RenderedDocument(output.getvalue(), DOCX_MIME, None, page_budget(data), False)
