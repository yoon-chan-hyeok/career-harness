import base64
import io
import zipfile

import pytest
from career.document_exports import render_docx, render_pdf
from career.document_uploads import UploadDecision, UploadRequest, confirm_upload, upload_preview
from career.editor_drafts import DraftField, DraftRequest, read_draft, save_draft
from career.models import Artifact, Evidence, Workspace
from career.provenance import artifact_provenance
from docx import Document
from pydantic import ValidationError
from pypdf import PdfReader
from sqlalchemy import select


def upload_body(filename, content):
    return UploadRequest(filename=filename, content_base64=base64.b64encode(content).decode())


def test_upload_preview_never_confirms_facts_and_registration_is_account_scoped(db):
    preview = upload_preview(db, "alice", upload_body("experience.txt", "직접 분석 보고서를 작성했습니다.".encode()))
    assert db.scalars(select(Evidence)).all() == []
    with pytest.raises(ValueError, match="직접 확인"):
        confirm_upload(db, "alice", preview["id"], UploadDecision(title="분석"))
    decision = UploadDecision(title="분석", confirmed=True)
    saved = confirm_upload(db, "alice", preview["id"], decision)
    assert saved["evidence"]["status"] == "UNVERIFIED"
    assert saved["evidence"]["data"]["ownership"] == "UNKNOWN"
    repeated = confirm_upload(db, "alice", preview["id"], decision)
    assert repeated["evidence"]["id"] == saved["evidence"]["id"]
    assert repeated["evidence"]["body"] == saved["evidence"]["body"]
    with pytest.raises(LookupError):
        confirm_upload(db, "bob", preview["id"], decision)
    with pytest.raises(ValueError, match="이미 등록"):
        confirm_upload(db, "alice", preview["id"], decision.model_copy(update={"kind": "EXPRESSION"}))


def test_application_upload_stays_expression_not_candidate_fact(db):
    preview = upload_preview(db, "alice", upload_body("old.txt", "과거 지원서에 쓴 문장입니다.".encode()))
    result = confirm_upload(db, "alice", preview["id"], UploadDecision(title="표현", kind="EXPRESSION", confirmed=True))
    assert result["evidence"]["kind"] == "EXPRESSION"
    assert result["evidence"]["status"] == "UNVERIFIED"


def test_pdf_docx_roundtrip_keeps_korean_and_reports_real_overflow(db):
    text = "문서 마지막 근거를 유지했습니다."
    value = {"sections": [{"title": "프로젝트 경험", "bullets": [{"text": text}]}], "page_budget": 2}
    pdf = render_pdf(value)
    assert pdf.pages == 1 and pdf.page_budget == 2 and not pdf.overflow
    preview = upload_preview(db, "alice", upload_body("resume.pdf", pdf.content))
    assert text in preview["text"] and preview["pages"][0]["number"] == 1
    docx = render_docx(value)
    document = Document(io.BytesIO(docx.content))
    assert text in "\n".join(p.text for p in document.paragraphs)
    assert docx.pages is None
    preview = upload_preview(db, "alice", upload_body("resume.docx", docx.content))
    assert text in preview["text"] and preview["pages"] == []
    overflow = render_pdf({"text": ("실제 수행한 분석 결과와 근거를 정리했습니다.\n" * 200) + "END-UNCHANGED", "page_budget": 1})
    assert overflow.pages is not None and overflow.pages > 1 and overflow.overflow
    assert "END-UNCHANGED" in "".join(p.extract_text() for p in PdfReader(io.BytesIO(overflow.content)).pages)


def test_docx_table_order_and_external_reference_rejection(db):
    document = Document()
    document.add_paragraph("표 앞의 문장")
    document.add_table(1, 1).cell(0, 0).text = "표 안의 근거"
    document.add_paragraph("표 뒤의 문장")
    stream = io.BytesIO()
    document.save(stream)
    value = upload_preview(db, "alice", upload_body("table.docx", stream.getvalue()))
    assert value["text"].index("표 앞") < value["text"].index("표 안") < value["text"].index("표 뒤")
    hostile = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(stream.getvalue())) as source, zipfile.ZipFile(hostile, "w") as target:
        for name in source.namelist():
            content = source.read(name)
            if name == "word/_rels/document.xml.rels":
                content = content.replace(b"</Relationships>", b"<Relationship Id='external' Type='image' Target='https://invalid.local/x' TargetMode = 'External'/></Relationships>")
            target.writestr(name, content)
    with pytest.raises(ValueError, match="외부 파일"):
        upload_preview(db, "alice", upload_body("external.docx", hostile.getvalue()))


@pytest.mark.parametrize("name,raw", [("a.docx", b"not zip"), ("a.pdf", b"%PDF-1.7\n/Type /X#52ef\nstartxref\n1\n%%EOF"), ("a.exe", b"abc"), ("a.txt", b"\0"), ("large.txt", b"a" * 2000001)], ids=["invalid-docx", "compressed-xref", "exe", "null", "oversize"])
def test_hostile_or_unsupported_upload_has_no_partial_import(db, name, raw):
    with pytest.raises(ValueError):
        upload_preview(db, "alice", upload_body(name, raw))
    assert db.scalars(select(Evidence)).all() == []


def test_recovery_draft_cannot_store_approval_secret_or_cross_user_base(db):
    ws = Workspace(user_id="alice", company="Demo", role="Analyst", jd_raw="Report")
    db.add(ws)
    db.flush()
    artifact = Artifact(user_id="alice", workspace_id=ws.id, kind="Draft", slot="q1", data={"text": "old"})
    db.add(artifact)
    db.flush()
    body = DraftRequest(payload={"0:text": DraftField(value="미저장 문장", kind="textarea")}, workspace_id=ws.id, base_artifact_id=artifact.id)
    save_draft(db, "alice", "editor:q1", body)
    assert read_draft(db, "alice", "editor:q1")["payload"]["0:text"]["value"] == "미저장 문장"
    with pytest.raises(LookupError):
        read_draft(db, "bob", "editor:q1")
    with pytest.raises(LookupError):
        save_draft(db, "bob", "editor:q1", body)
    for payload in ({"0:confirmed": {"value": "true", "kind": "text"}}, {"0:token": {"value": "private", "kind": "text"}}, {"0:text": {"value": "true", "kind": "checkbox"}}):
        with pytest.raises(ValidationError):
            DraftRequest.model_validate({"payload": payload})
    newer = Artifact(user_id="alice", workspace_id=ws.id, kind="Draft", slot="q1", version=2, data={"text": "new"})
    db.add(newer)
    db.flush()
    assert read_draft(db, "alice", "editor:q1")["stale_base"]
    with pytest.raises(ValueError, match="문서 버전"):
        save_draft(db, "alice", "editor:q1", body)


def test_provenance_uses_historical_research_dependency_and_no_other_users(db):
    research = Artifact(user_id="alice", workspace_id="w", kind="CompanyResearchResult", data={"sources": [{"id": "company", "title": "Old", "url": "https://example.com/old", "quote": "예전 원문"}]})
    private = Evidence(user_id="bob", source_id="bob-source", title="private", body="PRIVATE", data={})
    db.add_all([research, private])
    db.flush()
    newer = Artifact(user_id="alice", workspace_id="w", kind="CompanyResearchResult", version=2, data={"sources": [{"id": "company", "title": "New", "quote": "다른 원문"}]})
    document = Artifact(user_id="alice", workspace_id="w", kind="Draft", depends_on=[research.id], data={"claims": [{"text": "회사 설명", "evidence_ids": [private.id], "company_source_ids": ["company"]}]})
    db.add_all([newer, document])
    db.flush()
    result = artifact_provenance(db, "alice", document.id)
    assert result["items"][0]["company_sources"][0]["quote"] == "예전 원문"
    assert result["items"][0]["evidence"] == []
    assert "PRIVATE" not in str(result)
    with pytest.raises(LookupError):
        artifact_provenance(db, "bob", document.id)
